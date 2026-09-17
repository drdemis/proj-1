import os
import hashlib
import tkinter as tk
from tkinter import messagebox, simpledialog, ttk
from datetime import datetime
import openpyxl
from tkcalendar import DateEntry
from PIL import Image, ImageTk

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
EXCEL_FILE = os.path.join(BASE_DIR, "Lebanese_Army_Visits.xlsx")
LOGO_FILE = os.path.join(BASE_DIR, "defense_logo.png")
CEDAR_LOGO_FILE = os.path.join(BASE_DIR, "CEDAR_LOGO_FILE.png")
PASSWORD_FILE = os.path.join(BASE_DIR, "admin_password.hash")
DEVELOPER_PASSWORD_FILE = os.path.join(BASE_DIR, "developer_password.hash")

root = None
notebook = None
main_tab = None
developer_tab = None
cal_date = None
entry_visitor = None
txt_delegation = None
txt_news = None
suggestions = None
edit_menu = None


def password_hash(password):
    return hashlib.sha256(password.encode("utf-8")).hexdigest()


def authenticate_admin(root):
    if not os.path.exists(PASSWORD_FILE):
        while True:
            password = simpledialog.askstring(
                "إعداد كلمة مرور المسؤول",
                "أنشئ كلمة مرور للمسؤول:",
                parent=root,
                show="*"
            )
            if password is None:
                return False
            if len(password) < 4:
                messagebox.showwarning("كلمة مرور غير صالحة", "يجب أن تتكون كلمة المرور من 4 أحرف على الأقل.", parent=root)
                continue

            confirmation = simpledialog.askstring(
                "تأكيد كلمة المرور",
                "أعد إدخال كلمة المرور:",
                parent=root,
                show="*"
            )
            if password == confirmation:
                with open(PASSWORD_FILE, "w", encoding="ascii") as password_file:
                    password_file.write(password_hash(password))
                return True
            messagebox.showerror("خطأ", "كلمتا المرور غير متطابقتين.", parent=root)

    for attempt in range(3):
        password = simpledialog.askstring(
            "دخول المسؤول",
            "أدخل كلمة مرور المسؤول:",
            parent=root,
            show="*"
        )
        if password is None:
            return False
        with open(PASSWORD_FILE, "r", encoding="ascii") as password_file:
            saved_hash = password_file.read().strip()
        if password_hash(password) == saved_hash:
            return True
        messagebox.showerror("رفض الدخول", f"كلمة المرور غير صحيحة. المحاولة {attempt + 1} من 3.", parent=root)

    return False


def authenticate_developer(root):
    if not os.path.exists(DEVELOPER_PASSWORD_FILE):
        while True:
            password = simpledialog.askstring(
                "إعداد كلمة مرور المطور",
                "أنشئ كلمة مرور للمطور:",
                parent=root,
                show="*"
            )
            if password is None:
                return False
            if len(password) < 4:
                messagebox.showwarning("كلمة مرور غير صالحة", "يجب أن تتكون كلمة المرور من 4 أحرف على الأقل.", parent=root)
                continue

            confirmation = simpledialog.askstring(
                "تأكيد كلمة المرور",
                "أعد إدخال كلمة المرور:",
                parent=root,
                show="*"
            )
            if password == confirmation:
                with open(DEVELOPER_PASSWORD_FILE, "w", encoding="ascii") as password_file:
                    password_file.write(password_hash(password))
                return True
            messagebox.showerror("خطأ", "كلمتا المرور غير متطابقتين.", parent=root)

    password = simpledialog.askstring(
        "دخول المطور",
        "أدخل كلمة مرور المطور:",
        parent=root,
        show="*"
    )
    if password is None:
        return False
    with open(DEVELOPER_PASSWORD_FILE, "r", encoding="ascii") as password_file:
        saved_hash = password_file.read().strip()
    if password_hash(password) == saved_hash:
        return True
    messagebox.showerror("رفض الدخول", "كلمة مرور المطور غير صحيحة.", parent=root)
    return False


def select_all(widget):
    if isinstance(widget, tk.Text):
        widget.tag_add(tk.SEL, "1.0", tk.END)
        widget.mark_set(tk.INSERT, "1.0")
    else:
        widget.select_range(0, tk.END)
        widget.icursor(tk.END)
    return "break"


def show_edit_menu(event):
    widget = event.widget
    widget.focus_set()
    if edit_menu is None:
        return
    edit_menu.delete(0, tk.END)
    edit_menu.add_command(label="لصق", command=lambda: widget.event_generate("<<Paste>>"))
    edit_menu.add_command(label="نسخ", command=lambda: widget.event_generate("<<Copy>>"))
    edit_menu.add_command(label="قص", command=lambda: widget.event_generate("<<Cut>>"))
    edit_menu.add_separator()
    edit_menu.add_command(label="تحديد الكل", command=lambda: select_all(widget))
    edit_menu.tk_popup(event.x_root, event.y_root)


def enable_pasting(widget):
    widget.bind("<Button-3>", show_edit_menu)
    widget.bind("<Control-v>", lambda event: widget.event_generate("<<Paste>>"))
    widget.bind("<Control-V>", lambda event: widget.event_generate("<<Paste>>"))


def init_excel():
    if not os.path.exists(EXCEL_FILE):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "الزيارات"
        ws.append(["تاريخ الإدخال", "تاريخ الزيارة", "اسم الزائر", "الوفد المرافق", "الخبر / التفاصيل"])
        wb.save(EXCEL_FILE)


def save_data():
    if cal_date is None or entry_visitor is None or txt_delegation is None or txt_news is None:
        return False

    date_val = cal_date.get_date().strftime("%Y-%m-%d")
    visitor_val = entry_visitor.get().strip()
    delegation_val = txt_delegation.get("1.0", tk.END).strip()
    news_val = txt_news.get("1.0", tk.END).strip()

    if not visitor_val or not news_val:
        messagebox.showwarning("تنبيه", "يرجى إدخال اسم الزائر وتفاصيل الخبر على الأقل!")
        return False

    try:
        wb = openpyxl.load_workbook(EXCEL_FILE)
        ws = wb["الزيارات"]
        now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        ws.append([now_str, date_val, visitor_val, delegation_val, news_val])
        wb.save(EXCEL_FILE)
        messagebox.showinfo("نجاح", "تم حفظ البيانات في ملف Excel بنجاح!")
        clear_fields()
        return True
    except Exception as exc:
        messagebox.showerror("خطأ", f"حدث خطأ أثناء الحفظ: {exc}")
        return False


def search_data():
    if entry_visitor is None:
        return False

    query = entry_visitor.get().strip()
    if not query:
        messagebox.showwarning("تنبيه", "اكتب اسم الزائر في خانة 'اسم الزائر' للبحث عنه.")
        return False

    if not os.path.exists(EXCEL_FILE):
        messagebox.showinfo("بحث", "لا توجد سجلات محفوظة بعد.")
        return False

    wb = openpyxl.load_workbook(EXCEL_FILE)
    ws = wb["الزيارات"]
    results = []

    for row in ws.iter_rows(min_row=2, values_only=True):
        if query.lower() in str(row[2]).lower():
            results.append((row[1], row[2], row[3], row[4]))

    if results:
        if root is None:
            return False
        result_window = tk.Toplevel(root)
        result_window.title(f"نتائج البحث ({len(results)})")
        result_window.geometry("600x420")
        result_window.transient(root)

        results_frame = tk.Frame(result_window)
        results_frame.pack(fill="both", expand=True, padx=10, pady=10)

        for date_value, visitor_value, delegation_value, news_value in results:
            tk.Label(
                results_frame,
                text=f"التاريخ: {date_value}\nالزائر: {visitor_value}\nالوفد: {delegation_value}",
                font=("Arial", 11),
                anchor="e",
                justify="right"
            ).pack(fill="x", pady=(0, 3))

            news_text = tk.Text(results_frame, font=("Arial", 11), wrap="word", height=4, undo=False)
            news_text.pack(fill="x", pady=(0, 10))
            news_text.insert("1.0", str(news_value or ""))
            enable_pasting(news_text)
        return True

    messagebox.showinfo("بحث", "لم يتم العثور على نتائج مطابقة.")
    return False


def get_visitor_names():
    if not os.path.exists(EXCEL_FILE):
        return []

    try:
        wb = openpyxl.load_workbook(EXCEL_FILE, read_only=True)
        ws = wb["الزيارات"]
        names = {str(row[2]).strip() for row in ws.iter_rows(min_row=2, values_only=True) if row[2]}
        return sorted(names)
    except Exception:
        return []


def hide_suggestions(event=None):
    if suggestions is not None:
        suggestions.place_forget()


def choose_suggestion(event=None):
    if suggestions is None or entry_visitor is None:
        return
    selection = suggestions.curselection()
    if selection:
        entry_visitor.delete(0, tk.END)
        entry_visitor.insert(0, suggestions.get(selection[0]))
    hide_suggestions()
    entry_visitor.focus_set()


def update_suggestions(event=None):
    if entry_visitor is None or suggestions is None or root is None:
        return
    query = entry_visitor.get().strip().lower()
    if not query:
        hide_suggestions()
        return

    matches = [name for name in get_visitor_names() if query in name.lower()]
    suggestions.delete(0, tk.END)
    for name in matches:
        suggestions.insert(tk.END, name)

    if matches:
        root_x = entry_visitor.winfo_rootx() - root.winfo_rootx()
        root_y = entry_visitor.winfo_rooty() - root.winfo_rooty()
        suggestions.place(x=root_x, y=root_y + entry_visitor.winfo_height(), width=entry_visitor.winfo_width())
    else:
        hide_suggestions()


def clear_fields():
    if cal_date is None or entry_visitor is None or txt_delegation is None or txt_news is None:
        return
    cal_date.set_date(datetime.now())
    entry_visitor.delete(0, tk.END)
    txt_delegation.delete("1.0", tk.END)
    txt_news.delete("1.0", tk.END)


def show_developer_tab():
    if root is None or notebook is None or developer_tab is None:
        return
    if not authenticate_developer(root):
        return
    if str(developer_tab) not in notebook.tabs():
        notebook.add(developer_tab, text="المطور")
    notebook.select(developer_tab)


def main():
    global root, notebook, main_tab, developer_tab, cal_date, entry_visitor, txt_delegation, txt_news, suggestions, edit_menu

    init_excel()

    root = tk.Tk()
    root.title("دائرة العلاقات العامة والاعلام  - برنامج إدارة الزيارات والوفود")
    root.geometry("650x670")
    root.resizable(False, False)
    edit_menu = tk.Menu(root, tearoff=False)

    if not authenticate_admin(root):
        root.destroy()
        raise SystemExit

    notebook = ttk.Notebook(root)
    notebook.pack(fill="both", expand=True)
    main_tab = tk.Frame(notebook)
    developer_tab = tk.Frame(notebook)
    notebook.add(main_tab, text="الزيارات")

    canvas = tk.Canvas(main_tab, width=650, height=670, highlightthickness=0)
    canvas.pack(fill="both", expand=True)

    bg_photo = None
    if os.path.exists(LOGO_FILE):
        raw_img = Image.open(LOGO_FILE).convert("RGBA")
        raw_img = raw_img.resize((650, 670), Image.Resampling.LANCZOS)
        alpha = raw_img.split()[3]
        alpha = alpha.point(lambda p: int(p * 0.35))
        raw_img.putalpha(alpha)
        bg_photo = ImageTk.PhotoImage(raw_img)
        bg_item = canvas.create_image(0, 0, image=bg_photo, anchor="nw")
        canvas.tag_lower(bg_item)
    else:
        canvas.create_rectangle(0, 0, 650, 670, fill="#EAEAEA", outline="")

    canvas.create_rectangle(0, 0, 650, 95, fill="#1E5631", outline="")

    canvas.create_text(
        325, 24,
        text="الجمهورية اللبنانية",
        font=("Arial", 11, "bold"),
        fill="white",
        justify="center"
    )

    if os.path.exists(CEDAR_LOGO_FILE):
        cedar_img = Image.open(CEDAR_LOGO_FILE).convert("RGBA")
        cedar_img = cedar_img.resize((28, 28), Image.Resampling.LANCZOS)
        cedar_photo = ImageTk.PhotoImage(cedar_img)
        canvas.create_image(430, 24, image=cedar_photo, anchor="center")

    canvas.create_text(
        310, 52,
        text="وزارة الدفاع الوطني | دائرة العلاقات العامة والاعلام - برنامج إدارة الزيارات",
        font=("Arial", 11, "bold"),
        fill="white",
        justify="center"
    )

    if os.path.exists(LOGO_FILE):
        logo_img = Image.open(LOGO_FILE).convert("RGBA")
        logo_img = logo_img.resize((24, 24), Image.Resampling.LANCZOS)
        header_logo = ImageTk.PhotoImage(logo_img)
        canvas.create_image(605, 52, image=header_logo, anchor="center")

    label_bg = "#EAEAEA" if not bg_photo else "#FFFFFF"

    tk.Label(canvas, text=":التاريخ", font=("Arial", 11, "bold"), bg=label_bg, anchor="e").place(x=405, y=145, width=120, height=35)
    cal_date = DateEntry(canvas, width=28, background='#1E5631', foreground='white', date_pattern='yyyy-mm-dd', font=("Arial", 11), justify="right")
    cal_date.place(x=115, y=145, width=270, height=35)

    tk.Label(canvas, text=":اسم الزائر الكريم", font=("Arial", 11, "bold"), bg=label_bg, anchor="e").place(x=405, y=205, width=120, height=35)
    entry_visitor = tk.Entry(canvas, font=("Arial", 11), justify="right", width=30)
    entry_visitor.place(x=115, y=205, width=270, height=35)
    enable_pasting(entry_visitor)

    suggestions = tk.Listbox(root, font=("Arial", 10), height=4, justify="right", exportselection=False)
    entry_visitor.bind("<KeyRelease>", update_suggestions)
    entry_visitor.bind("<FocusOut>", lambda event: root.after(150, hide_suggestions))
    suggestions.bind("<ButtonRelease-1>", choose_suggestion)

    tk.Label(canvas, text=":الوفد المرافق", font=("Arial", 11, "bold"), bg=label_bg, anchor="e").place(x=405, y=265, width=120, height=35)
    txt_delegation = tk.Text(canvas, font=("Arial", 10), width=30, height=3)
    txt_delegation.place(x=115, y=265, width=270, height=70)
    enable_pasting(txt_delegation)

    tk.Label(canvas, text=":الخبر / التفاصيل", font=("Arial", 11, "bold"), bg=label_bg, anchor="e").place(x=405, y=355, width=120, height=35)
    txt_news = tk.Text(canvas, font=("Arial", 10), width=30, height=5)
    txt_news.place(x=115, y=355, width=270, height=105)
    enable_pasting(txt_news)

    btn_save = tk.Button(main_tab, text="+ إضافة زيارة جديدة", font=("Arial", 11, "bold"), bg="#1E5631", fg="white", width=18, command=save_data)
    canvas.create_window(470, 600, window=btn_save)

    btn_search = tk.Button(main_tab, text="🔍 بحث", font=("Arial", 11, "bold"), bg="#0056B3", fg="white", width=12, command=search_data)
    canvas.create_window(180, 600, window=btn_search)

    btn_developer = tk.Button(main_tab, text="دخول المطور", font=("Arial", 10, "bold"), command=show_developer_tab)
    canvas.create_window(325, 635, window=btn_developer)

    tk.Label(
        developer_tab,
        text="منطقة المطور - الوصول محمي بكلمة مرور",
        font=("Arial", 14, "bold")
    ).pack(pady=(35, 15))
    tk.Button(
        developer_tab,
        text="فتح ملف التطبيق للتعديل",
        font=("Arial", 11, "bold"),
        command=lambda: os.startfile(os.path.abspath(__file__))
    ).pack(pady=10)

    root.mainloop()


if __name__ == "__main__":
    main()
